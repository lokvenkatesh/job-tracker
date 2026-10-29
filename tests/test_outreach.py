import io
import json
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace

import docx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.db.models import OutreachCampaign, OutreachEmail, Resume
from app.main import create_app
from app.outreach import service
from app.outreach.checks import check_draft, opening_line, ungrounded_numbers, word_count
from app.outreach.generator import Draft, DraftGenerator, GenerationError
from app.outreach.recipients import Recipient, parse_recipients, recently_contacted
from app.outreach.resume import ResumeError, parse_resume
from tests.test_api import PASSWORD, settings

NOW = datetime(2026, 10, 29, 12, tzinfo=UTC)
RESUME_TEXT = (
    "Lok Vasamsetti\nMachine Learning Engineer\n"
    "Built a retrieval pipeline at Acme that cut p95 latency by 40% for 2M monthly users.\n"
    "Led a team of 4 to ship an LLM evaluation harness used across 12 product teams.\n"
    "Skills: Python, PyTorch, LangGraph, PostgreSQL, AWS.\n" + "Experience with production ML. " * 5
)
JD = "We are hiring a Machine Learning Engineer at Globex to build LLM agents. " * 3


# --- fixtures --------------------------------------------------------------


def make_pdf(lines: list[str]) -> bytes:
    """A minimal, valid single-page PDF with Helvetica text."""
    stream = (
        "BT /F1 10 Tf 40 760 Td 12 TL "
        + " ".join(
            "(" + ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") Tj T*"
            for ln in lines
        )
        + " ET"
    )
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = "%PDF-1.4\n", []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    return out.encode("latin-1")


def make_docx(paragraphs: list[str], table_rows=()) -> bytes:
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table_rows:
        t = d.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def good_body(name="Priya", company="Globex", extra="", opener=None):
    words = (
        f"Hi {name},\n\n{opener or f'I saw {company} is building LLM agents for logistics and the platform work caught my eye.'} "
        "At Acme I built a retrieval pipeline that cut p95 latency by 40% for 2M monthly users, and I led "
        "a team of 4 shipping an evaluation harness used by 12 product teams. That mix of shipping models "
        "and measuring them seems close to what your team needs right now. "
        f"{extra}Would you be open to a 15-minute chat next week about the Machine Learning Engineer role "
        "and what the team is focused on? Happy to work around your schedule, and no worries if the timing "
        "is wrong.\n\nThanks,\nLok"
    )
    return words


class FakeGenerator:
    def __init__(self, drafts=None, fail_brief=False):
        self.drafts = drafts or {}
        self.fail_brief = fail_brief
        self.calls = []

    def brief(self, jd, resume, notes):
        if self.fail_brief:
            raise GenerationError("Claude declined to write this one")
        return {
            "candidate_name": "Lok",
            "role": "Machine Learning Engineer",
            "company": "Globex",
            "top_skills": ["Python", "LLM agents", "PyTorch"],
            "matching_experience": ["Acme retrieval"],
            "wins": ["cut p95 latency by 40% for 2M users"],
        }

    def draft(self, brief, recipient, *, notes, sources, other_openings, instructions=None):
        self.calls.append((recipient.email, list(other_openings), instructions))
        result = self.drafts.get(recipient.email)
        if isinstance(result, Exception):
            raise result
        if result:
            return result
        company = recipient.company or "Globex"
        opener = [
            f"I saw {company} is building LLM agents for logistics and the platform work caught my eye.",
            f"Your team's writeup on how {company} evaluates agent quality stuck with me.",
            f"Reading about {company}'s move into routing optimization made me want to reach out.",
        ][len(self.calls) % 3]
        body = good_body(recipient.name or "there", company, opener=opener)
        return Draft(subject="ML Engineer at Globex", body=body, issues=[])


# --- resume parsing ---------------------------------------------------------


def test_parse_pdf():
    text, ctype = parse_resume("cv.pdf", make_pdf(RESUME_TEXT.splitlines()))
    assert ctype == "application/pdf"
    assert "retrieval pipeline at Acme" in text and "40%" in text


def test_parse_docx_with_tables():
    data = make_docx(RESUME_TEXT.splitlines(), table_rows=[("Skills", "Kubernetes, Terraform")])
    text, ctype = parse_resume("cv.docx", data)
    assert ctype.endswith("wordprocessingml.document")
    assert "12 product teams" in text and "Kubernetes, Terraform" in text


@pytest.mark.parametrize(
    "name,data,match",
    [
        ("cv.pdf", b"", "empty"),
        ("cv.txt", b"hello world", "PDF or a Word"),
        ("cv.docx", b"PK\x03\x04garbage", "damaged"),
        ("cv.pdf", b"%PDF-1.4 garbage", "damaged|no text|almost no text"),
        ("cv.pdf", b"%PDF" + b"0" * (5 * 1024 * 1024), "larger than 5 MB"),
    ],
)
def test_parse_resume_errors(name, data, match):
    with pytest.raises(ResumeError, match=match):
        parse_resume(name, data)


def test_parse_scanned_pdf_has_no_text():
    with pytest.raises(ResumeError, match="almost no text"):
        parse_resume("cv.pdf", make_pdf(["Lok"]))


# --- recipients --------------------------------------------------------------


def test_parse_recipients_formats():
    text = (
        "jane@acme.com, Jane Doe, Acme, Engineering Manager\n"
        "Bob Smith <BOB@Globex.com>\n"
        "carol@initech.io\n"
        "not-an-email\n"
        "jane@ACME.com\n"
        "\n"
    )
    ok, bad = parse_recipients(text)
    assert [r.email for r in ok] == ["jane@acme.com", "bob@globex.com", "carol@initech.io"]
    assert ok[0] == Recipient("jane@acme.com", "Jane Doe", "Acme", "Engineering Manager")
    assert ok[1].name == "Bob Smith"
    assert [(b.line, b.reason.split(":")[0]) for b in bad] == [
        (4, "invalid email"),
        (5, "duplicate in this list"),
    ]


def test_parse_recipients_csv_with_header_any_order():
    text = "Title,Email,Company,Name\nRecruiter,ana@ramp.com,Ramp,Ana Ortiz\n"
    ok, bad = parse_recipients(text)
    assert ok == [Recipient("ana@ramp.com", "Ana Ortiz", "Ramp", "Recruiter")] and bad == []


def test_parse_recipients_tsv():
    ok, _ = parse_recipients("email\tname\nx@y.com\tX Y\n")
    assert ok == [Recipient("x@y.com", "X Y")]


def test_recently_contacted(session):
    c = OutreachCampaign(jd_text="jd")
    session.add(c)
    session.flush()
    for email, status, sent_days_ago in [
        ("recent@x.com", "sent", 3),
        ("replied@x.com", "replied", 10),
        ("old@x.com", "sent", 45),
        ("draft@x.com", "draft", None),
    ]:
        session.add(
            OutreachEmail(
                campaign_id=c.id,
                recipient_email=email,
                subject="s",
                body="b",
                status=status,
                sent_at=NOW - timedelta(days=sent_days_ago) if sent_days_ago else None,
            )
        )
    session.flush()
    emails = ["recent@x.com", "replied@x.com", "old@x.com", "draft@x.com", "new@x.com"]
    assert set(recently_contacted(session, emails, now=NOW, days=30)) == {
        "recent@x.com",
        "replied@x.com",
    }
    assert recently_contacted(session, emails, now=NOW, days=0) == {}


# --- draft checks ------------------------------------------------------------

R = Recipient("priya@globex.com", "Priya Shah", "Globex", "Engineering Manager")
SOURCES = RESUME_TEXT + JD


def test_good_draft_passes():
    body = good_body()
    assert 90 <= word_count(body) <= 150
    assert check_draft("ML Engineer at Globex", body, recipient=R, sources=SOURCES) == []


@pytest.mark.parametrize(
    "mutate,expected",
    [
        (lambda b: b.replace("I saw", "I hope this email finds you well. I saw"), "banned phrase"),
        (lambda b: b.replace("I saw", "I’m excited to apply! I saw"), "banned phrase"),
        (lambda b: b.replace("Thanks,", "Thanks!!"), "exclamation"),
        (lambda b: b.replace("40%", "65%"), "numbers not found"),
        (lambda b: b.replace("Hi Priya", "Hi [Name]"), "placeholder"),
        (lambda b: b.replace("Priya", "there"), "by name"),
        (lambda b: b.replace("Globex", "your company"), "doesn't mention Globex"),
        (lambda b: "Hi Priya,\n\nShort note about Globex.", "words"),
        (
            lambda b: b.replace("I saw", "I am a passionate, results-driven self-starter. I saw"),
            "buzzwords",
        ),
    ],
)
def test_check_failures(mutate, expected):
    issues = check_draft("ML Engineer at Globex", mutate(good_body()), recipient=R, sources=SOURCES)
    assert any(expected in i for i in issues), issues


def test_subject_checks():
    assert any(
        "missing subject" in i for i in check_draft("", good_body(), recipient=R, sources=SOURCES)
    )
    long = "x" * 80
    assert any(
        "subject is 80" in i for i in check_draft(long, good_body(), recipient=R, sources=SOURCES)
    )


def test_opening_must_vary():
    body = good_body()
    other = opening_line(good_body(name="Sam"))
    issues = check_draft("s", body, recipient=R, sources=SOURCES, other_openings=[other])
    assert any("opening line" in i for i in issues)


def test_opening_line_skips_greeting():
    assert opening_line("Hi Priya,\n\nFirst sentence here. Second.") == "First sentence here."


def test_number_grounding_details():
    assert ungrounded_numbers("a 15-minute chat, or 20 min", "") == []
    assert ungrounded_numbers("2M users and 40% faster", "2M monthly users; 40% lower") == []
    assert ungrounded_numbers("grew revenue 3x", "nothing here") == ["3x"]


# --- generator (fake Anthropic client) --------------------------------------


class ScriptedAnthropic:
    def __init__(self, payloads):
        self.payloads = list(payloads)
        self.requests = []
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self._create))

    def _create(self, **kw):
        self.requests.append(kw)
        p = self.payloads.pop(0)
        if p == "refusal":
            return SimpleNamespace(stop_reason="refusal", content=[])
        return SimpleNamespace(
            stop_reason="end_turn", content=[SimpleNamespace(type="text", text=json.dumps(p))]
        )


def gen(payloads):
    client = ScriptedAnthropic(payloads)
    return DraftGenerator(client, settings(), model="claude-opus-5-5"), client


def test_brief_trims_lists():
    g, client = gen(
        [
            {
                "candidate_name": "Lok",
                "role": "MLE",
                "company": "Globex",
                "top_skills": ["a", "b", "c", "d", "e"],
                "matching_experience": ["x"],
                "wins": ["1", "2", "3"],
            }
        ]
    )
    b = g.brief(JD, RESUME_TEXT, None)
    assert len(b["top_skills"]) == 4 and len(b["wins"]) == 2
    assert client.requests[0]["output_config"]["format"]["type"] == "json_schema"
    assert "<resume>" in client.requests[0]["messages"][0]["content"]


def test_draft_retries_once_with_feedback():
    bad = {"subject": "Hello", "body": "Hi Priya,\n\nI hope this email finds you well. Globex."}
    good = {"subject": "ML Engineer at Globex", "body": good_body()}
    g, client = gen([bad, good])
    d = g.draft({"company": "Globex"}, R, notes=None, sources=SOURCES, other_openings=[])
    assert d.issues == [] and d.body == good["body"]
    assert len(client.requests) == 2
    retry_prompt = client.requests[1]["messages"][0]["content"]
    assert "Rewrite it to fix these problems" in retry_prompt and "banned phrase" in retry_prompt


def test_draft_keeps_remaining_issues_after_retry():
    bad = {"subject": "Hello", "body": "Hi Priya,\n\nShort. Globex."}
    g, client = gen([bad, bad])
    d = g.draft({}, R, notes=None, sources=SOURCES, other_openings=[])
    assert d.issues and len(client.requests) == 2


def test_draft_request_includes_openings_and_instructions():
    g, client = gen([{"subject": "ML Engineer at Globex", "body": good_body()}])
    g.draft(
        {"company": "Globex"},
        R,
        notes="I use their product",
        sources=SOURCES,
        other_openings=["Saw your talk on agents."],
        instructions="mention the talk",
    )
    content = client.requests[0]["messages"][0]["content"]
    assert "Saw your talk on agents." in content and "mention the talk" in content
    assert "I use their product" in content and "Engineering Manager" in content


def test_refusal_raises():
    g, _ = gen(["refusal"])
    with pytest.raises(GenerationError):
        g.brief(JD, RESUME_TEXT, None)


# --- service + API -------------------------------------------------------------


@pytest.fixture
def client(session_factory):
    app = create_app(settings(), session_factory)
    app.state.draft_generator = FakeGenerator()
    c = TestClient(app, follow_redirects=False)
    c.post("/auth/login", data={"username": "me", "password": PASSWORD})
    return c


@pytest.fixture
def resume_id(client):
    r = client.post(
        "/resumes",
        files={"file": ("cv.pdf", make_pdf(RESUME_TEXT.splitlines()), "application/pdf")},
    )
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_outreach_requires_login(session_factory):
    c = TestClient(create_app(settings(), session_factory))
    assert c.get("/outreach/campaigns").status_code == 401
    assert c.post("/resumes").status_code == 401


def test_resume_upload_and_default(client, session_factory):
    first = client.post(
        "/resumes", files={"file": ("a.pdf", make_pdf(RESUME_TEXT.splitlines()), "application/pdf")}
    ).json()
    second = client.post(
        "/resumes", files={"file": ("b.docx", make_docx(RESUME_TEXT.splitlines()), "x")}
    ).json()
    listed = client.get("/resumes").json()
    assert [r["id"] for r in listed] == [second["id"], first["id"]]
    assert listed[0]["is_default"] and not listed[1]["is_default"]
    assert "parsed_text" not in listed[0] and listed[0]["characters"] > 200
    with session_factory() as s:
        assert s.get(Resume, second["id"]).content.startswith(b"PK")
    bad = client.post("/resumes", files={"file": ("x.txt", b"hello", "text/plain")})
    assert bad.status_code == 422 and "PDF or a Word" in bad.json()["detail"]


def test_create_campaign_generates_drafts(client, resume_id):
    r = client.post(
        "/outreach/campaigns",
        json={
            "jd_text": JD,
            "recipients": "priya@globex.com, Priya Shah, Globex, Engineering Manager\nsam@globex.com\nbad-email",
            "notes": "I use Globex daily",
            "attach_resume": True,
        },
    )
    assert r.status_code == 201, r.text
    body = r.json()
    assert [x["reason"].split(":")[0] for x in body["rejected"]] == ["invalid email"]
    campaign = client.get(f"/outreach/campaigns/{body['id']}").json()  # background task already ran
    assert (campaign["role"], campaign["company"]) == ("Machine Learning Engineer", "Globex")
    emails = campaign["emails"]
    assert [e["recipient_email"] for e in emails] == ["priya@globex.com", "sam@globex.com"]
    assert all(e["status"] == "draft" and not e["pending"] and e["attach_resume"] for e in emails)
    assert 90 <= emails[0]["words"] <= 150
    gen_calls = client.app.state.draft_generator.calls
    assert gen_calls[0][1] == [] and len(gen_calls[1][1]) == 1  # second draft saw the first opening


def test_create_campaign_validation(client, resume_id):
    r = client.post("/outreach/campaigns", json={"jd_text": "too short", "recipients": "a@b.com"})
    assert r.status_code == 422 and "job description" in r.json()["detail"]
    r = client.post("/outreach/campaigns", json={"jd_text": JD, "recipients": "nope"})
    assert r.status_code == 422 and "No valid recipient" in r.json()["detail"]


def test_create_campaign_needs_resume(client):
    r = client.post("/outreach/campaigns", json={"jd_text": JD, "recipients": "a@b.com"})
    assert r.status_code == 422 and "resume" in r.json()["detail"]


def test_recent_recipients_skipped_unless_override(client, resume_id, session_factory):
    with session_factory() as s:
        c = OutreachCampaign(jd_text="old")
        s.add(c)
        s.flush()
        s.add(
            OutreachEmail(
                campaign_id=c.id,
                recipient_email="priya@globex.com",
                subject="s",
                body="b",
                status="sent",
                sent_at=datetime.now(UTC) - timedelta(days=5),
            )
        )
        s.commit()
    r = client.post(
        "/outreach/campaigns",
        json={"jd_text": JD, "recipients": "priya@globex.com\nsam@globex.com"},
    ).json()
    emails = client.get(f"/outreach/campaigns/{r['id']}").json()["emails"]
    assert emails[0]["status"] == "skipped" and "within 30 days" in emails[0]["skip_reason"]
    assert emails[0]["body"] == "" and emails[1]["status"] == "draft"
    r = client.post(
        "/outreach/campaigns",
        json={"jd_text": JD, "recipients": "priya@globex.com", "override_recent": True},
    ).json()
    assert client.get(f"/outreach/campaigns/{r['id']}").json()["emails"][0]["status"] == "draft"


def test_generation_failure_is_recorded(client, resume_id):
    client.app.state.draft_generator = FakeGenerator(
        drafts={"sam@globex.com": GenerationError("cut off")}
    )
    r = client.post(
        "/outreach/campaigns",
        json={"jd_text": JD, "recipients": "priya@globex.com\nsam@globex.com"},
    ).json()
    emails = client.get(f"/outreach/campaigns/{r['id']}").json()["emails"]
    assert emails[0]["issues"] == []
    assert emails[1]["issues"] == ["Generation failed: cut off"] and not emails[1]["pending"]


def test_brief_failure_marks_all_failed(client, resume_id):
    client.app.state.draft_generator = FakeGenerator(fail_brief=True)
    r = client.post(
        "/outreach/campaigns", json={"jd_text": JD, "recipients": "a@b.com\nc@d.com"}
    ).json()
    emails = client.get(f"/outreach/campaigns/{r['id']}").json()["emails"]
    assert all(e["issues"][0].startswith("Generation failed") for e in emails)


def _campaign_with_two(client):
    r = client.post(
        "/outreach/campaigns",
        json={"jd_text": JD, "recipients": "priya@globex.com, Priya Shah, Globex\nsam@globex.com"},
    ).json()
    return client.get(f"/outreach/campaigns/{r['id']}").json()


def test_patch_and_approval_rules(client, resume_id):
    c = _campaign_with_two(client)
    e = c["emails"][0]
    r = client.patch(f"/outreach/emails/{e['id']}", json={"status": "approved"})
    assert r.json()["status"] == "approved"
    # editing approved text drops it back to draft and re-runs the checks
    r = client.patch(
        f"/outreach/emails/{e['id']}", json={"body": e["body"].replace("40%", "90%")}
    ).json()
    assert r["status"] == "draft" and any("numbers not found" in i for i in r["issues"])
    # edit + approve in one request keeps it approved
    r = client.patch(
        f"/outreach/emails/{e['id']}", json={"body": e["body"], "status": "approved"}
    ).json()
    assert r["status"] == "approved" and r["issues"] == []
    assert client.patch(f"/outreach/emails/{e['id']}", json={"status": "sent"}).status_code == 409
    assert client.patch(f"/outreach/emails/{e['id']}", json={"bogus": 1}).status_code == 422
    r = client.patch(f"/outreach/emails/{e['id']}", json={"subject": "", "status": "approved"})
    assert r.status_code == 409 and "subject" in r.json()["detail"]


def test_sent_emails_are_immutable(client, resume_id, session_factory):
    c = _campaign_with_two(client)
    eid = c["emails"][0]["id"]
    with session_factory() as s:
        s.get(OutreachEmail, eid).status = "sent"
        s.commit()
    assert client.patch(f"/outreach/emails/{eid}", json={"body": "x"}).status_code == 409
    assert client.post(f"/outreach/emails/{eid}/regenerate").status_code == 409


def test_regenerate_with_instructions(client, resume_id):
    c = _campaign_with_two(client)
    eid = c["emails"][1]["id"]
    client.patch(f"/outreach/emails/{eid}", json={"status": "approved"})
    r = client.post(f"/outreach/emails/{eid}/regenerate", json={"instructions": "shorter"}).json()
    assert r["status"] == "draft"  # a new version always needs a fresh approval
    assert client.app.state.draft_generator.calls[-1][2] == "shorter"


def test_approve_all_holds_back_drafts_with_issues(client, resume_id):
    client.app.state.draft_generator = FakeGenerator(
        drafts={"sam@globex.com": Draft("s", "Hi Sam, short.", ["12 words (aim for 90-150)"])}
    )
    c = _campaign_with_two(client)
    assert client.post(f"/outreach/campaigns/{c['id']}/approve-all").json() == {
        "approved": 1,
        "held_back": 1,
    }
    statuses = [e["status"] for e in client.get(f"/outreach/campaigns/{c['id']}").json()["emails"]]
    assert statuses == ["approved", "draft"]
    assert client.post("/outreach/campaigns/999/approve-all").status_code == 404


def test_list_campaigns(client, resume_id):
    c = _campaign_with_two(client)
    client.patch(f"/outreach/emails/{c['emails'][0]['id']}", json={"status": "approved"})
    listed = client.get("/outreach/campaigns").json()
    assert listed[0]["id"] == c["id"] and listed[0]["counts"] == {"approved": 1, "draft": 1}


def test_nothing_is_sent(client, resume_id, session_factory):
    c = _campaign_with_two(client)
    client.post(f"/outreach/campaigns/{c['id']}/approve-all")
    with session_factory() as s:
        assert all(
            e.sent_at is None and e.gmail_message_id is None
            for e in s.scalars(select(OutreachEmail))
        )


# --- pages -------------------------------------------------------------------


def test_pages_flow(client, resume_id):
    html = client.get("/dashboard/outreach").text
    assert "New campaign" in html and "cv.pdf" in html and "Cold mail" in html
    r = client.post(
        "/dashboard/outreach",
        data={
            "jd_text": JD,
            "recipients": "priya@globex.com, Priya Shah, Globex\nnope",
            "resume_id": str(resume_id),
        },
    )
    assert r.status_code == 303 and r.headers["location"].endswith("?rejected=1")
    page = client.get(r.headers["location"]).text
    assert "Priya Shah" in page and "1 line skipped" in page and "Approve all 1 ready" in page
    assert (
        "What the emails draw on" in page and 'hx-trigger="every 3s"' not in page
    )  # nothing pending


def test_page_form_errors_keep_input(client):
    r = client.post(
        "/dashboard/outreach", data={"jd_text": JD, "recipients": "nope@", "notes": "keep me"}
    )
    assert r.status_code == 422 and "keep me" in r.text and "No valid recipient" in r.text


def test_page_draft_actions(client, resume_id):
    c = _campaign_with_two(client)
    e = c["emails"][0]
    hx = {"HX-Request": "true"}
    r = client.post(
        f"/dashboard/outreach/emails/{e['id']}",
        data={"action": "approve", "subject": e["subject"], "body": e["body"]},
        headers=hx,
    )
    assert r.status_code == 200 and "pill-approved" in r.text and "Unapprove" in r.text
    r = client.post(f"/dashboard/outreach/emails/{e['id']}", data={"action": "skip"}, headers=hx)
    assert "pill-skipped" in r.text
    r = client.post(
        f"/dashboard/outreach/emails/{e['id']}/regenerate",
        data={"instructions": "warmer"},
        headers=hx,
    )
    assert "pill-draft" in r.text
    r = client.post(f"/dashboard/outreach/{c['id']}/approve-all", headers=hx)
    assert "Approved 2." in r.text
    assert (
        client.post(f"/dashboard/outreach/emails/{e['id']}", data={"action": "send"}).status_code
        == 422
    )


def test_pending_drafts_poll(client, resume_id, session_factory):
    c = _campaign_with_two(client)
    with session_factory() as s:
        s.get(OutreachEmail, c["emails"][1]["id"]).generated_at = None
        s.commit()
    page = client.get(f"/dashboard/outreach/{c['id']}").text
    assert 'hx-trigger="every 3s"' in page and "Writing 1 draft" in page


def test_service_save_resume_switches_default(session):
    a = service.save_resume(session, "a.pdf", make_pdf(RESUME_TEXT.splitlines()))
    b = service.save_resume(session, "b.pdf", make_pdf(RESUME_TEXT.splitlines()))
    assert not a.is_default and b.is_default and service.default_resume(session).id == b.id

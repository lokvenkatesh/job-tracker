"""Resume upload parsing: PDF (pdfplumber) and DOCX (python-docx) to plain text."""

import io
import re
import zipfile

MAX_BYTES = 5 * 1024 * 1024
PDF_TYPE = "application/pdf"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ResumeError(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace(" ", " ").replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _pdf_text(data: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:  # many resume templates lay out sections in tables
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def parse_resume(filename: str, data: bytes) -> tuple[str, str]:
    """Returns (text, content_type). The file type is checked by content, not just name."""
    if not data:
        raise ResumeError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ResumeError("The file is larger than 5 MB.")
    name = filename.lower()
    try:
        if data.startswith(b"%PDF"):
            text, content_type = _pdf_text(data), PDF_TYPE
        elif data.startswith(b"PK") and name.endswith(".docx"):
            text, content_type = _docx_text(data), DOCX_TYPE
        else:
            raise ResumeError("Upload a PDF or a Word (.docx) file.")
    except ResumeError:
        raise
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as e:
        raise ResumeError("Couldn't read that file; it may be damaged.") from e
    except Exception as e:  # pdfminer raises a variety of parser errors
        raise ResumeError("Couldn't read that file; it may be damaged.") from e
    text = _clean(text)
    if len(text) < 200:
        raise ResumeError(
            "Found almost no text in this file. If it's a scanned image, export a text-based PDF."
        )
    return text, content_type
